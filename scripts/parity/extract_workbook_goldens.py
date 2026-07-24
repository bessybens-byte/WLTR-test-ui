#!/usr/bin/env python3
"""Extract MS4 VOC 230120 preload goldens + MassHunter paste text from the unlocked workbook."""

from __future__ import annotations

import json
import re
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

WORKBOOK = Path(
    r"c:\Users\AbdulaiAwal\Downloads\wltr\docs\230120_MS4_VOC_Preloaded (5) - Unlocked.xlsm"
)
OUT_DIR = Path(__file__).resolve().parents[2] / "fixtures" / "ms4-voc-230120"
RUNS_DIR = OUT_DIR / "runs"

# Cal Data paste bands (column index 1-based), 20 slots in VBA clear list
CAL_BAND_COLS = [1, 11, 21, 31, 41, 51, 61, 71, 81, 91, 101, 111, 121]

# Excel DVD model codes → WLTR workbook variant set
MODEL_CODE_MAP = {
    "A": {"regressionType": "Average", "weightingMode": "None", "label": "Average RF"},
    "LS": {"regressionType": "Linear", "weightingMode": "None", "label": "LS (EW)"},
    "LSIC": {"regressionType": "Linear", "weightingMode": "InverseX", "label": "LS (InvX)"},
    "LSISC": {
        "regressionType": "Linear",
        "weightingMode": "InverseXSquared",
        "label": "LS (InvX2)",
    },
    "LSF0": {
        "regressionType": "LinearForcedZero",
        "weightingMode": "None",
        "label": "LS Forced 0",
    },
    "Q": {"regressionType": "Quadratic", "weightingMode": "None", "label": "Quad (EW)"},
    "QIC": {
        "regressionType": "Quadratic",
        "weightingMode": "InverseX",
        "label": "Quad (InvX)",
    },
    "QISC": {
        "regressionType": "Quadratic",
        "weightingMode": "InverseXSquared",
        "label": "Quad (InvX2)",
    },
    "QF0": {
        "regressionType": "QuadraticForcedZero",
        "weightingMode": "None",
        "label": "Quad Forced 0",
    },
}

LEVEL_LADDER = [
    {"sample": "01 VOC CAL", "trueConcentration": 0.1, "sortOrder": 1, "file": "0120-06.D"},
    {"sample": "02 VOC CAL", "trueConcentration": 0.2, "sortOrder": 2, "file": "0120-07.D"},
    {"sample": "05 VOC CAL", "trueConcentration": 0.5, "sortOrder": 3, "file": "0120-08.D"},
    {"sample": "1 VOC CAL", "trueConcentration": 1.0, "sortOrder": 4, "file": "0120-09.D"},
    {"sample": "2 VOC CAL", "trueConcentration": 2.0, "sortOrder": 5, "file": "0120-10.D"},
    {"sample": "4 VOC CAL", "trueConcentration": 4.0, "sortOrder": 6, "file": "0120-11.D"},
    {"sample": "8 VOC CAL", "trueConcentration": 8.0, "sortOrder": 7, "file": "0120-12.D"},
    {"sample": "20 VOC CAL", "trueConcentration": 20.0, "sortOrder": 8, "file": "0120-13.D"},
    {"sample": "50 VOC CAL", "trueConcentration": 50.0, "sortOrder": 9, "file": "0120-14.D"},
    {"sample": "100 VOC CAL", "trueConcentration": 100.0, "sortOrder": 10, "file": "0120-15.D"},
    {"sample": "120 VOC CAL", "trueConcentration": 120.0, "sortOrder": 11, "file": "0120-16.D"},
    {"sample": "140 VOC CAL", "trueConcentration": 140.0, "sortOrder": 12, "file": "0120-17.D"},
    {"sample": "160 VOC CAL", "trueConcentration": 160.0, "sortOrder": 13, "file": "0120-22.D"},
]


def cell_str(v) -> str | None:
    if v is None:
        return None
    if isinstance(v, float):
        return str(v)
    return str(v).strip()


def extract_column_text(ws, col: int, max_row: int = 250) -> str:
    lines: list[str] = []
    trailing_blank = 0
    for r in range(1, max_row + 1):
        v = ws.cell(r, col).value
        if v is None:
            trailing_blank += 1
            if trailing_blank > 8 and lines:
                break
            lines.append("")
            continue
        trailing_blank = 0
        lines.append(str(v).rstrip())
    # trim trailing empties
    while lines and not lines[-1].strip():
        lines.pop()
    return "\n".join(lines) + ("\n" if lines else "")


def parse_sample_name(raw: str) -> str | None:
    m = re.search(r"Sample\s*:\s*(.+)", raw, re.I)
    if not m:
        return None
    return m.group(1).strip()


def parse_acq_on(raw: str) -> str | None:
    m = re.search(r"Acq On\s*:\s*(.+)", raw, re.I)
    if not m:
        return None
    return m.group(1).strip()


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    RUNS_DIR.mkdir(parents=True, exist_ok=True)

    print("Loading workbook (values)…")
    wb_v = load_workbook(WORKBOOK, data_only=True, keep_vba=False)
    print("Loading workbook (formulas)…")
    wb_f = load_workbook(WORKBOOK, data_only=False, keep_vba=False)

    ref = wb_v["Ref Table"]
    sr = wb_v["Summary Report"]
    cal = wb_v["Cal Data"]
    icv = wb_v["ICV Data"]

    # --- Config from Ref Table ---
    config = {
        "instrumentId": cell_str(ref["Q12"].value),
        "method": cell_str(ref["Q15"].value),
        "analyst": cell_str(ref["Q18"].value),
        "peer": cell_str(ref["Q21"].value),
        "unitsSwitch": ref["R27"].value,
        "rsdPercentLimit": float(ref["Q40"].value),
        "minCorrelationR": float(ref["Q43"].value),
        "minCorrelationR2": float(ref["R44"].value),
        "labelModeSwitch": int(ref["R47"].value),  # 2 = r²
        "icalPctDiff": float(ref["Q52"].value),
        "pctDiffHighFactor": float(ref["Q53"].value),
        "pctDiffLowFactor": float(ref["S53"].value),
        "maxMissedPoints": int(ref["Q56"].value),
        "maxRSE": float(ref["Q59"].value),
        "minPointsAverage": int(ref["Q64"].value),
        "minPointsLS": int(ref["R64"].value),
        "minPointsQuad": int(ref["S64"].value),
        "icvCdsParityPercent": float(ref["Q70"].value),
        "icvLimitPercent": float(ref["Q77"].value),
        "soilDilutionFactor": float(ref["Q81"].value),
        "aqueousDilutionFactor": float(ref["Q87"].value),
        # Effective UI bounds: Excel factors 0.8/1.2 on ±20% → -16 / +24
        "pctDiffLowBound": -float(ref["Q52"].value) * float(ref["S53"].value),
        "pctDiffHighBound": float(ref["Q52"].value) * float(ref["Q53"].value),
        "isRsdPercentLimit": 20.0,  # Summary Report M14 / P14 (documented)
        "labelMode": "RSquared" if int(ref["R47"].value) == 2 else "R",
        "quantitationMode": "InternalStandard",
        "minPointsRequired": int(ref["S64"].value),  # strictest (Quad=6)
        "levelLadder": LEVEL_LADDER,
        "modelCodeMap": MODEL_CODE_MAP,
    }

    # Internal standards / surrogates
    internal_standards = []
    for r in range(21, 24):
        name = cell_str(ref.cell(r, 23).value)  # W
        conc = ref.cell(r, 24).value  # X
        on = ref.cell(r, 21).value  # U
        if name and on == 1:
            internal_standards.append({"name": name, "concentration": float(conc), "role": "InternalStandard"})

    surrogates = []
    for r in range(34, 37):
        name = cell_str(ref.cell(r, 23).value)
        conc = ref.cell(r, 24).value
        on = ref.cell(r, 21).value
        if name and on == 1:
            surrogates.append(
                {
                    "name": name,
                    "concentration": float(conc) if conc is not None else 10.0,
                    "role": "Surrogate",
                    "surrogateSpikeAmount": float(conc) if conc is not None else 10.0,
                    "surrogateRecoveryLowerLimit": 70.0,
                    "surrogateRecoveryUpperLimit": 130.0,
                }
            )

    # Per-analyte Ref Table criteria (from row 46 until blank)
    analyte_criteria = []
    for r in range(46, 200):
        on = ref.cell(r, 21).value  # U
        name = cell_str(ref.cell(r, 23).value)  # W
        if not name:
            break
        if on != 1:
            continue
        mb_aq = ref.cell(r, 24).value  # X
        spcc = ref.cell(r, 26).value  # Z Min RF
        ccc = ref.cell(r, 27).value  # AA Min %RSD
        icv_true = ref.cell(r, 29).value  # AC
        lcl = ref.cell(r, 30).value  # AD
        ucl = ref.cell(r, 31).value  # AE
        analyte_criteria.append(
            {
                "name": name,
                "methodBlankLimit": float(mb_aq) if mb_aq is not None else None,
                "isSpcc": spcc is not None and spcc != "",
                "minResponseFactor": float(spcc) if spcc not in (None, "") else None,
                "isCcc": ccc is not None and ccc != "",
                "maxRsdPercent": float(ccc) if ccc not in (None, "") else None,
                "icvLcsConcentration": float(icv_true) if icv_true is not None else None,
                "icvLcsLowerControlLimit": float(lcl) if lcl is not None else None,
                "icvLcsUpperControlLimit": float(ucl) if ucl is not None else None,
                "concentrationMultiplier": 1.0,
            }
        )

    # --- Summary Report Page 2 analytes (skip blank gaps between sections) ---
    summary_analytes = []
    blank_streak = 0
    for r in range(57, 200):
        name = cell_str(sr.cell(r, 3).value)
        ical = cell_str(sr.cell(r, 30).value)
        if not name:
            blank_streak += 1
            if blank_streak > 8 and summary_analytes:
                break
            continue
        blank_streak = 0
        if ical not in ("Acceptable", "Unacceptable", "Fail", "Failed"):
            continue
        model_code = cell_str(sr.cell(r, 9).value)
        if model_code not in MODEL_CODE_MAP:
            # chosen model must be a known DVD code
            continue
        summary_analytes.append(
            {
                "row": r,
                "name": name,
                "rlAqueous": cell_str(sr.cell(r, 4).value),
                "rlSoil": cell_str(sr.cell(r, 5).value),
                "averageRf": sr.cell(r, 6).value,
                "rsdPercent": sr.cell(r, 7).value,
                "rsdPass": cell_str(sr.cell(r, 8).value),
                "chosenModelCode": model_code,
                "chosenModel": MODEL_CODE_MAP.get(model_code or "", None),
                "calModelPass": cell_str(sr.cell(r, 10).value),
                "rSquared": sr.cell(r, 11).value,
                "rSquaredPass": cell_str(sr.cell(r, 12).value),
                "spccOrCcc": cell_str(sr.cell(r, 13).value),
                "totalCalPoints": sr.cell(r, 14).value,
                "minCalPoints": sr.cell(r, 15).value,
                "metMinPoints": cell_str(sr.cell(r, 16).value),
                "pointsWithinRefit": sr.cell(r, 17).value,
                "failures": sr.cell(r, 18).value,
                "failuresPass": cell_str(sr.cell(r, 19).value),
                "rsePass": cell_str(sr.cell(r, 20).value),
                "icalSummary": ical,
            }
        )

    # --- IS pairing from DVD Internal Standard Response vs all CAL IS areas ---
    dvd = wb_v["Data Visualization Deck (DVD)"]

    # Build IS area lookup: area -> is_name from every CAL paste
    area_to_is: dict[float, str] = {}
    for col in CAL_BAND_COLS:
        text = extract_column_text(cal, col)
        if "Quantitation Report" not in text:
            continue
        for line in text.splitlines():
            m = re.search(
                r"\)\s+(Fluorobenzene|Chlorobenzene-d5|1,4-Dichlorobenzene-d4)\s+"
                r"[\d.]+\s+\d+\s+(\d+)",
                line,
            )
            if m:
                area_to_is[float(m.group(2))] = m.group(1)

    cal1_is = {
        "Fluorobenzene": None,
        "Chlorobenzene-d5": None,
        "1,4-Dichlorobenzene-d4": None,
    }
    for area, name in area_to_is.items():
        # prefer first-seen (cal1) for reporting
        if cal1_is[name] is None:
            cal1_is[name] = area

    is_by_analyte: dict[str, str] = {}
    for r in range(1, min(dvd.max_row or 1, 2800)):
        label = cell_str(dvd.cell(r, 4).value)
        if label != "Internal Standard Response":
            continue
        name = cell_str(dvd.cell(r, 2).value)
        if not name:
            continue
        votes: dict[str, int] = {}
        for c in range(6, 30):
            v = dvd.cell(r, c).value
            if not isinstance(v, (int, float)):
                continue
            is_name = area_to_is.get(float(v))
            if is_name:
                votes[is_name] = votes.get(is_name, 0) + 1
        if votes:
            is_by_analyte[name] = max(votes.items(), key=lambda kv: kv[1])[0]

    # --- Upload fixtures ---
    runs_manifest = []
    for idx, col in enumerate(CAL_BAND_COLS):
        raw = extract_column_text(cal, col)
        if not raw.strip() or "Quantitation Report" not in raw:
            continue
        sample = parse_sample_name(raw) or LEVEL_LADDER[idx]["sample"]
        level_meta = next((L for L in LEVEL_LADDER if L["sample"] == sample), None)
        if level_meta is None and idx < len(LEVEL_LADDER):
            level_meta = LEVEL_LADDER[idx]
            sample = level_meta["sample"]
        fname = f"cal_{idx+1:02d}_{re.sub(r'[^A-Za-z0-9]+', '_', sample).strip('_')}.txt"
        (RUNS_DIR / fname).write_text(raw, encoding="utf-8")
        runs_manifest.append(
            {
                "runType": "CAL",
                "file": f"runs/{fname}",
                "sample": sample,
                "level": sample,
                "trueConcentration": level_meta["trueConcentration"] if level_meta else None,
                "sortOrder": level_meta["sortOrder"] if level_meta else idx + 1,
                "acqOn": parse_acq_on(raw),
                "bandColumn": get_column_letter(col),
            }
        )

    icv_raw = extract_column_text(icv, 1)
    icv_name = "icv_VOC_ICV1.txt"
    (RUNS_DIR / icv_name).write_text(icv_raw, encoding="utf-8")
    runs_manifest.append(
        {
            "runType": "ICV",
            "file": f"runs/{icv_name}",
            "sample": parse_sample_name(icv_raw) or "VOC ICV1",
            "level": None,
            "acqOn": parse_acq_on(icv_raw),
            "bandColumn": "A",
        }
    )

    # Unique model codes used
    codes = sorted({a["chosenModelCode"] for a in summary_analytes if a["chosenModelCode"]})

    # Attach default IS onto criteria / summary where known
    for row in analyte_criteria:
        row["defaultInternalStandard"] = is_by_analyte.get(row["name"])
    for row in summary_analytes:
        row["defaultInternalStandard"] = is_by_analyte.get(row["name"])

    goldens = {
        "workbook": str(WORKBOOK),
        "config": config,
        "internalStandards": internal_standards,
        "surrogates": surrogates,
        "analyteCriteria": analyte_criteria,
        "summaryAnalytes": summary_analytes,
        "isAssignments": is_by_analyte,
        "cal1InternalStandardAreas": cal1_is,
        "runs": runs_manifest,
        "modelCodesSeen": codes,
        "counts": {
            "summaryAnalytes": len(summary_analytes),
            "analyteCriteriaOn": len(analyte_criteria),
            "isAssignments": len(is_by_analyte),
            "calRuns": sum(1 for r in runs_manifest if r["runType"] == "CAL"),
            "icvRuns": sum(1 for r in runs_manifest if r["runType"] == "ICV"),
        },
    }

    out_path = OUT_DIR / "goldens.json"
    out_path.write_text(json.dumps(goldens, indent=2, default=str), encoding="utf-8")
    print(f"Wrote {out_path}")
    print(json.dumps(goldens["counts"], indent=2))
    print("Model codes seen:", codes)
    print("pctDiff bounds:", config["pctDiffLowBound"], config["pctDiffHighBound"])

    wb_v.close()
    wb_f.close()


if __name__ == "__main__":
    main()
